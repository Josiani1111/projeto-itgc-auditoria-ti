
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from src.controle_acesso import verificar_acessos, usuarios


def gerar_relatorio_excel():
    alertas = verificar_acessos(usuarios)

    pasta_relatorios = (
        Path(__file__).resolve().parent.parent / "relatorios"
    )
    pasta_relatorios.mkdir(parents=True, exist_ok=True)

    arquivo = pasta_relatorios / "relatorio_auditoria.xlsx"

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Auditoria ITGC"

    planilha.append([
        "Data da auditoria",
        "Controle",
        "Funcionário",
        "Falha identificada",
        "Risco",
        "Ação recomendada",
        "Status",
    ])

    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    data_auditoria = datetime.now().strftime("%d/%m/%Y %H:%M")

    if alertas:
        for alerta in alertas:
            nome = alerta.split(":")[1].split(" está")[0].strip()

            planilha.append([
                data_auditoria,
                "Gestão de acessos",
                nome,
                alerta,
                "Acesso não autorizado a sistemas e dados",
                "Solicitar a desativação do acesso e confirmar a execução",
                "Pendente de verificação",
            ])
    else:
        planilha.append([
            data_auditoria,
            "Gestão de acessos",
            "N/A",
            "Nenhuma falha identificada",
            "Nenhum risco identificado nesta verificação",
            "Manter os controles e realizar verificações periódicas",
            "Sem falhas identificadas",
        ])

    larguras = {
        "A": 20,
        "B": 24,
        "C": 18,
        "D": 55,
        "E": 40,
        "F": 55,
        "G": 25,
    }

    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura

    for linha in planilha.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions

    workbook.save(arquivo)

    print(f"Relatório Excel gerado em: {arquivo}")


if __name__ == "__main__":
    gerar_relatorio_excel()
